"""
campaign.py — Cruise/campaign management logic (GUI-free).

Ports the processing core of two former standalone tkinter tools so the PyQt6
"Cruise" menu can drive them natively:

  * File-registry + line-coordinates Excel generation (from FicherosExcelTopas):
    :func:`build_registro_excel` and :func:`build_coordenadas_excel`.
  * Acquisition statistics — ping rate, ping interval and vessel speed from the
    SEG-Y time/navigation headers (from PingRateTopas):
    :func:`calculate_metrics_for_phase`.

GUI-independence contract (see core/__init__): no tkinter/PyQt here. ``openpyxl``
is heavy and only needed for the Excel builders, so it is imported lazily inside
those functions; ``segyio``/``pyproj`` are already core dependencies.

Directory layout expected by every entry point
-----------------------------------------------
A base directory that either:
  * contains ``SGY/`` and/or ``RAW/`` subfolders, each holding one folder per
    seismic line (``SGY/<line>/*.sgy``, ``RAW/<line>/*.raw``), or
  * is itself the parent of the per-line folders (no SGY/RAW split).
:func:`resolve_bases` detects which layout is in use.
"""
from __future__ import annotations

import math
import re
from datetime import datetime, timedelta
from pathlib import Path
from typing import Callable, List, Optional, Tuple

import pyproj
import segyio

# ── Line-name ordering ──────────────────────────────────────────────────────────

def sort_key(name: str) -> tuple:
    """TL < L for the same number; number ascending; suffix A→B→C."""
    m = re.match(r"^(TL|L)(\d+)([A-Z]*)$", name.upper())
    if not m:
        return (9999, 1, "")
    prefix, number, suffix = m.groups()
    return (int(number), 0 if prefix == "TL" else 1, suffix)


def resolve_bases(base_dir: Path) -> Tuple[Path, Path]:
    """Detect the folder layout, returning ``(sgy_base, raw_base)``.

    Layout A: ``base_dir/SGY/<line>/`` and ``base_dir/RAW/<line>/``.
    Layout B: ``base_dir/<line>/`` (the base dir is itself the lines root).
    """
    sgy_sub = base_dir / "SGY"
    raw_sub = base_dir / "RAW"
    if sgy_sub.exists() or raw_sub.exists():
        return sgy_sub, raw_sub
    return base_dir, base_dir


def gather_line_folders(sgy_base: Path, raw_base: Path) -> list:
    folders: set = set()
    for d in (sgy_base, raw_base):
        if d.exists():
            folders.update(f.name for f in d.iterdir() if f.is_dir())
    return sorted(folders, key=sort_key)


# ── SEG-Y header readers (shared) ────────────────────────────────────────────────

def sgy_trace_coords(filepath: Path, trace_idx: int) -> Tuple[Optional[float], Optional[float]]:
    """Read ``(lon, lat)`` in decimal degrees from one trace of a SEG-Y file.

    ``trace_idx`` = 0 for the first trace, -1 for the last. Applies the
    SourceGroupScalar and converts arc-seconds (CoordinateUnits == 2) to degrees.
    Returns ``(None, None)`` on any read error or an all-zero coordinate.
    """
    try:
        with segyio.open(str(filepath), ignore_geometry=True) as f:
            h = f.header[trace_idx]
            scalar = h[segyio.TraceField.SourceGroupScalar] or 1
            cu = h[segyio.TraceField.CoordinateUnits]
            div = abs(scalar) if scalar < 0 else (1.0 / scalar if scalar > 0 else 1)
            x = h[segyio.TraceField.SourceX] / div
            y = h[segyio.TraceField.SourceY] / div
            if cu == 2:
                x /= 3600.0
                y /= 3600.0
            if x == 0 and y == 0:
                return None, None
            return float(x), float(y)
    except Exception:
        return None, None


def datetime_from_header(header) -> Optional[datetime]:
    """Convert SEG-Y time-header bytes (157-166) to a ``datetime`` (or None)."""
    year = header[segyio.TraceField.YearDataRecorded]
    day = header[segyio.TraceField.DayOfYear]
    hour = header[segyio.TraceField.HourOfDay]
    minute = header[segyio.TraceField.MinuteOfHour]
    sec = header[segyio.TraceField.SecondOfMinute]
    if year == 0 and day == 0:
        return None
    if year < 100:
        year += 2000
    try:
        return datetime(year, 1, 1) + timedelta(
            days=day - 1, hours=hour, minutes=minute, seconds=sec)
    except ValueError:
        return None


def coords_from_header(header) -> Tuple[float, float, int]:
    """Extract ``(x, y, coordinate_units)`` applying the scalar (bytes 71-80, 89)."""
    scalar = header[segyio.TraceField.SourceGroupScalar] or 1
    cu = header[segyio.TraceField.CoordinateUnits]
    div = abs(scalar) if scalar < 0 else (1.0 / scalar if scalar > 0 else 1)
    x = header[segyio.TraceField.SourceX] / div
    y = header[segyio.TraceField.SourceY] / div
    if cu == 2:
        x /= 3600.0
        y /= 3600.0
    return float(x), float(y), cu


# ── UTM projection helpers ───────────────────────────────────────────────────────

_proj_cache: dict = {}


def _get_transformer(zone: int, north: bool = True):
    epsg = 32600 + zone if north else 32700 + zone
    if epsg not in _proj_cache:
        _proj_cache[epsg] = pyproj.Transformer.from_crs(
            "EPSG:4326", f"EPSG:{epsg}", always_xy=True)
    return _proj_cache[epsg]


def auto_zone(lon: float) -> int:
    return math.floor((lon + 180) / 6) + 1


def to_utm(lon: Optional[float], lat: Optional[float],
           force_zone: Optional[int] = None):
    """Convert ``(lon, lat)`` WGS84 → ``(easting, northing, zone)``.

    Excel convention downstream: Y = easting, X = northing. ``force_zone``
    projects to a fixed zone so a line's start and end share a zone even when it
    crosses a UTM boundary (keeps the Euclidean length valid).
    """
    if lon is None or lat is None:
        return None, None, None
    try:
        zone = force_zone if force_zone else auto_zone(lon)
        transformer = _get_transformer(zone, north=(lat >= 0))
        easting, northing = transformer.transform(lon, lat)
        return easting, northing, zone
    except Exception:
        return None, None, None


# ════════════════════════════════════════════════════════════════════════════════
# File-registry Excel
# ════════════════════════════════════════════════════════════════════════════════

_COL_W_REG = [12, 8, 11, 8, 11, 24, 24, 24, 24, 10, 10, 20]


def _parse_dt(filename: str) -> Optional[datetime]:
    m = re.match(r"(\d{14})", Path(filename).stem)
    if not m:
        return None
    s = m.group(1)
    try:
        return datetime(int(s[0:4]), int(s[4:6]), int(s[6:8]),
                        int(s[8:10]), int(s[10:12]), int(s[12:14]))
    except ValueError:
        return None


def _excel_date(dt: datetime) -> int:
    return (dt - datetime(1899, 12, 30)).days


def _excel_time(dt: datetime) -> float:
    return (dt.hour * 3600 + dt.minute * 60 + dt.second) / 86400.0


def _get_files(dirs: list, ext: str) -> list:
    found: list = []
    for d in dirs:
        try:
            found += [f for f in Path(d).iterdir()
                      if f.is_file() and f.suffix.lower() == ext]
        except (FileNotFoundError, PermissionError):
            pass
    return sorted(set(found), key=lambda f: f.name)


def build_registro_excel(phases: list, project_name: str, output_path: str) -> List[str]:
    """Generate the file-registry Excel.

    ``phases``: list of dicts ``{"sheet_name", "label", "base_dir"}``.
    Returns a per-phase summary (one string per phase).
    """
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    yellow = PatternFill("solid", fgColor="FFFF99")
    blue = PatternFill("solid", fgColor="99CCFF")
    ctr = Alignment(horizontal="center", vertical="center")

    def bold(size=10):
        return Font(name="Arial", bold=True, size=size)

    def norm():
        return Font(name="Arial", bold=False, size=10)

    def write_headers(ws, phase_label):
        ws["A1"] = project_name
        ws["A1"].font = bold(12)
        ws["B2"] = "INICIO LINEA"; ws.merge_cells("B2:C2")
        ws["D2"] = "FINAL LINEA"; ws.merge_cells("D2:E2")
        for addr, fill in (("B2", yellow), ("D2", blue)):
            ws[addr].font = bold(); ws[addr].fill = fill; ws[addr].alignment = ctr
        headers = [
            "LINEA ID", "HORA", "FECHA", "HORA", "FECHA",
            "FICHERO INICIO (*.raw)", "FICHERO FINAL (*.raw)",
            "FICHERO INICIO (*.sgy)", "FICHERO FINAL (*.sgy)",
            "FIX INICIO", "FIX FINAL", "",
        ]
        hfills = [None, yellow, yellow, blue, blue,
                  yellow, blue, yellow, blue, yellow, blue, None]
        for col, (h, f) in enumerate(zip(headers, hfills), 1):
            cell = ws.cell(row=3, column=col, value=h)
            cell.font = bold(); cell.alignment = ctr
            if f:
                cell.fill = f
        ws.cell(row=4, column=1, value=phase_label).font = bold()

    def write_line(ws, row, line_name, sgy_dir, raw_dir):
        dirs = [sgy_dir, raw_dir]
        raw_files = _get_files(dirs, ".raw")
        sgy_files = _get_files(dirs, ".sgy")
        all_files = sorted(raw_files + sgy_files, key=lambda f: f.name)
        start_dt = _parse_dt(all_files[0].name) if all_files else None
        end_dt = _parse_dt(all_files[-1].name) if all_files else None
        ws.cell(row=row, column=1, value=line_name).font = norm()
        if start_dt:
            c = ws.cell(row=row, column=2, value=_excel_time(start_dt))
            c.number_format = "h:mm"; c.fill = yellow; c.alignment = ctr
            c = ws.cell(row=row, column=3, value=_excel_date(start_dt))
            c.number_format = "dd/mm/yy"; c.fill = yellow; c.alignment = ctr
        if end_dt:
            c = ws.cell(row=row, column=4, value=_excel_time(end_dt))
            c.number_format = "h:mm"; c.fill = blue; c.alignment = ctr
            c = ws.cell(row=row, column=5, value=_excel_date(end_dt))
            c.number_format = "dd/mm/yy"; c.fill = blue; c.alignment = ctr
        file_vals = [
            raw_files[0].name if raw_files else "",
            raw_files[-1].name if raw_files else "",
            sgy_files[0].name if sgy_files else "",
            sgy_files[-1].name if sgy_files else "",
        ]
        for col, val in enumerate(file_vals, 6):
            ws.cell(row=row, column=col, value=val).font = norm()

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    summary: List[str] = []

    for phase in phases:
        sheet_name = phase["sheet_name"]
        label = phase["label"]
        base_dir = Path(phase["base_dir"])
        sgy_base, raw_base = resolve_bases(base_dir)
        sorted_folders = gather_line_folders(sgy_base, raw_base)

        ws = wb.create_sheet(title=sheet_name)
        write_headers(ws, label)
        for i, w in enumerate(_COL_W_REG, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        for row_idx, folder_name in enumerate(sorted_folders, start=5):
            write_line(ws, row_idx, folder_name,
                       sgy_base / folder_name, raw_base / folder_name)
        summary.append(f"'{sheet_name}': {len(sorted_folders)} líneas")

    wb.save(output_path)
    return summary


# ════════════════════════════════════════════════════════════════════════════════
# Coordinates + line-length Excel
# ════════════════════════════════════════════════════════════════════════════════

_FMT_GEO = "0.0000000"
_FMT_UTM = "0.0000"
_FMT_LEN = "0.00"


def _get_line_coords(sgy_dir: Path, raw_dir: Path):
    """Return ``(lon_i, lat_i, lon_f, lat_f)`` for a line — first trace of the
    first file, last trace of the last file."""
    files: list = []
    for d in (sgy_dir, raw_dir):
        try:
            files += [f for f in Path(d).iterdir()
                      if f.is_file() and f.suffix.lower() == ".sgy"]
        except (FileNotFoundError, PermissionError):
            pass
    files = sorted(set(files), key=lambda f: f.name)
    if not files:
        return None, None, None, None
    lon_i, lat_i = sgy_trace_coords(files[0], 0)
    lon_f, lat_f = sgy_trace_coords(files[-1], -1)
    return lon_i, lat_i, lon_f, lat_f


def build_coordenadas_excel(phases: list, project_name: str, output_path: str) -> List[str]:
    """Generate the coordinates + line-length Excel.

    ``phases``: list of dicts ``{"label", "base_dir", "forced_zone"}``.
    Returns a per-phase summary (one string per phase).
    """
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    yellow_hdr = PatternFill("solid", fgColor="FFFF00")
    ctr = Alignment(horizontal="center", vertical="center")
    left = Alignment(horizontal="left", vertical="center")

    def bold(size=9):
        return Font(name="Arial", bold=True, size=size)

    def norm(size=9):
        return Font(name="Arial", bold=False, size=size)

    def write_coord_headers(ws):
        headers = ["Nombre", "Longitud inicial", "Latitud inicial",
                   "Longitud final", "Latitud final"]
        for c, h in enumerate(headers, 1):
            cell = ws.cell(row=1, column=c, value=h)
            cell.font = bold(); cell.alignment = ctr

    def write_longitud_headers(ws):
        headers = ["Nombre", "Longitud inicial", "Latitud inicial",
                   "Longitud final", "Latitud final",
                   "Y inicial (UTM)", "X inicial (UTM)",
                   "Y final (UTM)", "X final (UTM)",
                   "longitud (m)", "longitud (km)", "longitud (millas)"]
        for c, h in enumerate(headers, 1):
            cell = ws.cell(row=1, column=c, value=h)
            cell.font = bold(); cell.fill = yellow_hdr; cell.alignment = ctr

    def write_section(ws, row, label):
        cell = ws.cell(row=row, column=1, value=label)
        cell.font = bold(); cell.alignment = left

    def write_data(ws_coord, ws_long, row, name,
                   lon_i, lat_i, lon_f, lat_f, Y_i, X_i, Y_f, X_f):
        ws_coord.cell(row=row, column=1, value=name).font = bold()
        for col, val, fmt in ((2, lon_i, _FMT_GEO), (3, lat_i, _FMT_GEO),
                              (4, lon_f, _FMT_GEO), (5, lat_f, _FMT_GEO)):
            if val is not None:
                c = ws_coord.cell(row=row, column=col, value=val)
                c.number_format = fmt; c.font = norm()
        ws_long.cell(row=row, column=1, value=name).font = bold()
        for col, val, fmt in ((2, lon_i, _FMT_GEO), (3, lat_i, _FMT_GEO),
                              (4, lon_f, _FMT_GEO), (5, lat_f, _FMT_GEO),
                              (6, Y_i, _FMT_UTM), (7, X_i, _FMT_UTM),
                              (8, Y_f, _FMT_UTM), (9, X_f, _FMT_UTM)):
            if val is not None:
                c = ws_long.cell(row=row, column=col, value=val)
                c.number_format = fmt; c.font = norm()
        if all(v is not None for v in (Y_i, X_i, Y_f, X_f)):
            r = row
            c_m = ws_long.cell(row=r, column=10,
                               value=f"=SQRT((H{r}-F{r})^2+(I{r}-G{r})^2)")
            c_m.number_format = _FMT_LEN; c_m.font = norm()
            c_km = ws_long.cell(row=r, column=11, value=f"=J{r}/1000")
            c_km.number_format = _FMT_LEN; c_km.font = norm()
            c_nm = ws_long.cell(row=r, column=12, value=f"=J{r}/1852")
            c_nm.number_format = _FMT_LEN; c_nm.font = norm()

    def set_widths(ws, widths):
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    ws_coord = wb.create_sheet("COORDENADAS LINEAS")
    ws_long = wb.create_sheet("LONGITUD LINEAS")
    write_coord_headers(ws_coord)
    write_longitud_headers(ws_long)
    set_widths(ws_coord, [14, 14, 14, 14, 14])
    set_widths(ws_long, [14, 14, 14, 14, 14, 14, 14, 14, 14, 13, 13, 16])

    current_row = 2
    summary: List[str] = []

    for phase in phases:
        label = phase["label"]
        base_dir = Path(phase["base_dir"])
        user_zone = phase.get("forced_zone")
        sgy_base, raw_base = resolve_bases(base_dir)
        sorted_folders = gather_line_folders(sgy_base, raw_base)

        phase_zones: list = []
        section_row = current_row
        write_section(ws_coord, section_row, label)
        write_section(ws_long, section_row, label)
        current_row += 1

        n_ok = 0
        for folder_name in sorted_folders:
            sgy_dir = sgy_base / folder_name
            raw_dir = raw_base / folder_name
            lon_i, lat_i, lon_f, lat_f = _get_line_coords(sgy_dir, raw_dir)
            Y_i = X_i = Y_f = X_f = None
            used_zone = None
            if lon_i is not None and lat_i is not None:
                used_zone = user_zone if user_zone else auto_zone(lon_i)
                Y_i, X_i, _ = to_utm(lon_i, lat_i, force_zone=used_zone)
                phase_zones.append(used_zone)
            if lon_f is not None and lat_f is not None:
                Y_f, X_f, _ = to_utm(lon_f, lat_f, force_zone=used_zone)
            write_data(ws_coord, ws_long, current_row, folder_name,
                       lon_i, lat_i, lon_f, lat_f, Y_i, X_i, Y_f, X_f)
            if lon_i is not None:
                n_ok += 1
            current_row += 1

        if phase_zones:
            unique_zones = sorted(set(phase_zones))
            zones_str = "/".join(str(z) for z in unique_zones)
            forced_marker = " *" if user_zone else ""
            huso_label = f"{label} (HUSO {zones_str}{forced_marker})"
            ws_coord.cell(row=section_row, column=1).value = huso_label
            ws_long.cell(row=section_row, column=1).value = huso_label

        auto_str = "" if user_zone else " (auto)"
        summary.append(
            f"'{label}': {len(sorted_folders)} líneas, {n_ok} con coordenadas, "
            f"huso(s): {sorted(set(phase_zones))}{auto_str}"
            if phase_zones else
            f"'{label}': {len(sorted_folders)} líneas, sin coordenadas SGY")

    wb.save(output_path)
    return summary


def detect_zone_for_base(base_dir: str) -> Tuple[Optional[int], Optional[str], Optional[float]]:
    """Scan the first available ``.sgy`` under ``base_dir`` and return the
    ``(zone, filename, lon)`` of its first trace, or ``(None, None, None)`` if no
    readable SGY coordinate is found. Used by the UTM-zone detect button.
    """
    base = Path(base_dir)
    sgy_base, raw_base = resolve_bases(base)
    sgy_file: Optional[Path] = None
    for sub in (sgy_base, raw_base):
        if not sub.exists():
            continue
        for line_dir in sorted(sub.iterdir()):
            if not line_dir.is_dir():
                continue
            for f in sorted(line_dir.iterdir()):
                if f.suffix.lower() == ".sgy":
                    sgy_file = f
                    break
            if sgy_file:
                break
        if sgy_file:
            break
    if not sgy_file:
        return None, None, None
    lon, _lat = sgy_trace_coords(sgy_file, 0)
    if lon is None:
        return None, sgy_file.name, None
    return auto_zone(lon), sgy_file.name, lon


# ════════════════════════════════════════════════════════════════════════════════
# Acquisition statistics — ping rate / interval / vessel speed
# ════════════════════════════════════════════════════════════════════════════════

def calc_distance(x1: float, y1: float, x2: float, y2: float, cu: int) -> float:
    """Distance in metres — Haversine for geographic degrees (cu == 2), plain
    Euclidean for projected metres otherwise."""
    if cu == 2:
        R = 6371000.0
        phi1, phi2 = math.radians(y1), math.radians(y2)
        dphi = math.radians(y2 - y1)
        dlam = math.radians(x2 - x1)
        a = (math.sin(dphi / 2) ** 2
             + math.cos(phi1) * math.cos(phi2) * math.sin(dlam / 2) ** 2)
        c = 2 * math.atan2(math.sqrt(a), math.sqrt(1 - a))
        return R * c
    return math.sqrt((x2 - x1) ** 2 + (y2 - y1) ** 2)


def _resolve_sgy_base(base_dir: Path) -> Path:
    """Ping-rate variant: prefer a ``SGY/`` subfolder, else the dir itself."""
    sgy_sub = base_dir / "SGY"
    return sgy_sub if sgy_sub.exists() else base_dir


def calculate_metrics_for_phase(phase_name: str, base_dir_path: str):
    """Compute ping rate (Hz), ping interval (s) and vessel speed (knots) by
    iterating the SEG-Y files of a phase.

    Returns ``(log_text, rates_hz, intervals_s, speeds_knots)`` — the log is a
    human-readable per-line/per-phase report; the three lists are the raw
    per-file values (for a caller-side project-wide average).
    """
    base = Path(base_dir_path)
    sgy_base = _resolve_sgy_base(base)

    line_dirs = [d for d in sgy_base.iterdir() if d.is_dir()]
    if not line_dirs:
        line_dirs = [sgy_base]

    rates_hz_fase: List[float] = []
    intervals_s_fase: List[float] = []
    speeds_knots_fase: List[float] = []
    log_output = [f"\n{'=' * 70}\n📌 FASE: {phase_name}\n"
                  f"📂 Directorio: {base_dir_path}\n{'=' * 70}"]

    for line_dir in sorted(line_dirs):
        sgy_files = sorted(line_dir.glob("*.sgy"))
        if not sgy_files:
            continue
        log_output.append(f"\n  〰️ Línea: {line_dir.name}")
        rates_hz_linea: List[float] = []
        intervals_s_linea: List[float] = []
        speeds_knots_linea: List[float] = []

        for f in sgy_files:
            try:
                with segyio.open(str(f), ignore_geometry=True) as segy:
                    n_traces = segy.tracecount
                    if n_traces < 2:
                        log_output.append(
                            f"    [!] Ignorado {f.name}: Insuficientes trazas.")
                        continue
                    dt_start = datetime_from_header(segy.header[0])
                    dt_end = datetime_from_header(segy.header[n_traces - 1])
                    x_start, y_start, cu_start = coords_from_header(segy.header[0])
                    x_end, y_end, _cu_end = coords_from_header(segy.header[n_traces - 1])
                    if not dt_start or not dt_end:
                        log_output.append(
                            f"    [!] Ignorado {f.name}: Fechas a cero o inválidas.")
                        continue
                    delta_sec = (dt_end - dt_start).total_seconds()
                    if delta_sec > 0:
                        rate_hz = n_traces / delta_sec
                        interval_s = delta_sec / n_traces
                        rates_hz_linea.append(rate_hz)
                        intervals_s_linea.append(interval_s)
                        rates_hz_fase.append(rate_hz)
                        intervals_s_fase.append(interval_s)
                        nav_valida = ((x_start != 0 or y_start != 0)
                                      and (x_end != 0 or y_end != 0))
                        speed_str = "N/A"
                        if nav_valida:
                            dist_m = calc_distance(x_start, y_start, x_end, y_end, cu_start)
                            speed_knots = (dist_m / delta_sec) * 1.94384
                            speeds_knots_linea.append(speed_knots)
                            speeds_knots_fase.append(speed_knots)
                            speed_str = f"{speed_knots:.2f} nudos"
                        log_output.append(
                            f"    📄 {f.name}: {rate_hz:.2f} Hz | "
                            f"{interval_s:.3f} s | V: {speed_str}")
                    else:
                        log_output.append(
                            f"    [!] Ignorado {f.name}: Delta de tiempo es 0s.")
            except Exception as e:
                log_output.append(f"    [X] Error leyendo {f.name}: {e}")

        if rates_hz_linea:
            avg_hz = sum(rates_hz_linea) / len(rates_hz_linea)
            avg_s = sum(intervals_s_linea) / len(intervals_s_linea)
            avg_v = (sum(speeds_knots_linea) / len(speeds_knots_linea)
                     if speeds_knots_linea else 0.0)
            v_str = f"| Vel: {avg_v:.2f} nudos" if speeds_knots_linea else "| Vel: N/A"
            log_output.append(
                f"  > Media {line_dir.name}: {avg_hz:.2f} Hz | {avg_s:.3f} s {v_str}")

    if rates_hz_fase:
        avg_hz = sum(rates_hz_fase) / len(rates_hz_fase)
        avg_s = sum(intervals_s_fase) / len(intervals_s_fase)
        avg_v = (sum(speeds_knots_fase) / len(speeds_knots_fase)
                 if speeds_knots_fase else 0.0)
        v_str = (f"| Velocidad: {avg_v:.2f} nudos" if speeds_knots_fase
                 else "| Velocidad: Sin datos de NAV")
        log_output.append(
            f"\n📊 MEDIA DE LA FASE '{phase_name}': "
            f"{avg_hz:.2f} Hz | {avg_s:.3f} s {v_str}")
    else:
        log_output.append(
            f"\n⚠️ No se pudo calcular ninguna métrica para la fase '{phase_name}'.")

    return "\n".join(log_output), rates_hz_fase, intervals_s_fase, speeds_knots_fase
