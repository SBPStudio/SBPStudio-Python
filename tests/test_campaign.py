"""
test_campaign.py — headless verification of core/campaign.py.

Covers the geophysical math ported from the two former standalone tkinter
tools (FicherosExcelTopas, PingRateTopas) so it stays stable under future
refactors: line-name ordering, folder-layout detection, SEG-Y coordinate/time
header decoding, Haversine vs. Euclidean distance, UTM zone auto-detection,
and the two Excel builders (registry + coordinates/length).

No GUI here — sbp_studio.core.campaign is GUI-free by contract, and none of
these tests touch Qt.
"""
from __future__ import annotations

import math
from datetime import datetime
from pathlib import Path

import numpy as np
import openpyxl
import pytest
import segyio

from sbp_studio.core import campaign as C


# ── Local SEG-Y writer — full control over per-trace time/coords ──────────────
# tests/make_synthetic_segy.py's helper hardcodes "+1 minute per trace" and
# offers no way to place an exact coordinate delta, so the ping-rate/speed
# tests (which need EXACT, hand-picked deltas to assert exact Hz/knots values)
# use this minimal writer instead. The file-registry/coordinates tests reuse
# make_synthetic_segy since they only need trace *count* and endpoint
# coordinates, not exact time deltas.

def _write_metrics_sgy(path: str, *, xs, ys, seconds, cu: int = 1,
                       scalar: int = 1, year: int = 2024, doy: int = 100,
                       ns: int = 16) -> None:
    """Write a SEG-Y file with exact per-trace (x, y, second-of-minute).

    ``scalar=1`` (positive) makes ``coords_from_header``'s ``div`` exactly
    1.0, so the stored SourceX/SourceY equal ``xs``/``ys`` verbatim — no
    scalar-decoding subtlety to account for in the expected values.
    ``seconds`` must each be < 60 (single-minute window, no hour/day rollover)
    and monotonically non-decreasing.
    """
    n = len(xs)
    assert len(ys) == n and len(seconds) == n
    spec = segyio.spec()
    spec.sorting = None
    spec.format = 1
    spec.samples = list(range(ns))
    spec.tracecount = n
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    with segyio.create(path, spec) as f:
        f.bin.update(tsort=segyio.TraceSortingFormat.UNKNOWN_SORTING, hdt=1000, dto=1000)
        for i in range(n):
            f.header[i] = {
                segyio.TraceField.SourceX: int(round(xs[i])),
                segyio.TraceField.SourceY: int(round(ys[i])),
                segyio.TraceField.SourceGroupScalar: scalar,
                segyio.TraceField.CoordinateUnits: cu,
                segyio.TraceField.YearDataRecorded: year,
                segyio.TraceField.DayOfYear: doy,
                segyio.TraceField.HourOfDay: 10,
                segyio.TraceField.MinuteOfHour: 0,
                segyio.TraceField.SecondOfMinute: int(seconds[i]),
            }
            f.trace[i] = np.zeros(ns, dtype=np.float32)


# ════════════════════════════════════════════════════════════════════════════════
# Line-name ordering + folder-layout detection
# ════════════════════════════════════════════════════════════════════════════════

class TestSortKeyAndLayout:
    def test_tl_sorts_before_l_for_same_number(self):
        assert C.sort_key("TL5") < C.sort_key("L5")

    def test_numeric_ascending_not_lexicographic(self):
        assert C.sort_key("L2") < C.sort_key("L10")

    def test_suffix_ascending(self):
        assert C.sort_key("L3A") < C.sort_key("L3B") < C.sort_key("L3C")

    def test_unrecognised_name_falls_back(self):
        # No crash, and it sorts after any well-formed TL/L name.
        assert C.sort_key("MISC") > C.sort_key("L999Z")

    def test_resolve_bases_layout_a_sgy_raw_subfolders(self, tmp_path):
        (tmp_path / "SGY").mkdir()
        (tmp_path / "RAW").mkdir()
        sgy_base, raw_base = C.resolve_bases(tmp_path)
        assert sgy_base == tmp_path / "SGY"
        assert raw_base == tmp_path / "RAW"

    def test_resolve_bases_layout_b_flat_dir_is_lines_root(self, tmp_path):
        (tmp_path / "L1").mkdir()
        sgy_base, raw_base = C.resolve_bases(tmp_path)
        assert sgy_base == tmp_path
        assert raw_base == tmp_path

    def test_resolve_bases_sgy_only_still_layout_a(self, tmp_path):
        (tmp_path / "SGY").mkdir()
        sgy_base, raw_base = C.resolve_bases(tmp_path)
        assert sgy_base == tmp_path / "SGY"
        assert raw_base == tmp_path / "RAW"   # doesn't exist, but path returned

    def test_gather_line_folders_merges_sgy_and_raw_sorted(self, tmp_path):
        sgy = tmp_path / "SGY"; raw = tmp_path / "RAW"
        (sgy / "L2").mkdir(parents=True)
        (sgy / "TL1").mkdir()
        (raw / "L1").mkdir(parents=True)
        (raw / "L2").mkdir()   # overlaps with SGY/L2 — must dedupe
        folders = C.gather_line_folders(sgy, raw)
        assert folders == ["TL1", "L1", "L2"]

    def test_gather_line_folders_missing_dirs_returns_empty(self, tmp_path):
        assert C.gather_line_folders(tmp_path / "nope1", tmp_path / "nope2") == []


# ════════════════════════════════════════════════════════════════════════════════
# SEG-Y header decoding (coordinates + time)
# ════════════════════════════════════════════════════════════════════════════════

class TestHeaderDecoding:
    def test_sgy_trace_coords_arc_seconds_conversion(self, tmp_path):
        from tests.make_synthetic_segy import make_synthetic_segy
        p = make_synthetic_segy(str(tmp_path / "a.sgy"), n_traces=10, ns=32,
                                coord_unit=2, scalar_coord=-100,
                                base_lon=-8.0, base_lat=43.0, lon_step=0.001)
        lon0, lat0 = C.sgy_trace_coords(Path(p), 0)
        assert lon0 == pytest.approx(-8.0, abs=1e-6)
        assert lat0 == pytest.approx(43.0, abs=1e-6)
        lon_last, lat_last = C.sgy_trace_coords(Path(p), -1)
        assert lon_last == pytest.approx(-8.0 + 9 * 0.001, abs=1e-6)

    def test_sgy_trace_coords_all_zero_returns_none(self, tmp_path):
        _write_metrics_sgy(str(tmp_path / "z.sgy"), xs=[0, 0], ys=[0, 0], seconds=[0, 1])
        lon, lat = C.sgy_trace_coords(Path(tmp_path / "z.sgy"), 0)
        assert lon is None and lat is None

    def test_sgy_trace_coords_unreadable_file_returns_none(self, tmp_path):
        bogus = tmp_path / "not_a_segy.sgy"
        bogus.write_bytes(b"not a real segy file")
        lon, lat = C.sgy_trace_coords(bogus, 0)
        assert lon is None and lat is None

    def test_datetime_from_header_normal(self):
        h = {segyio.TraceField.YearDataRecorded: 2024,
             segyio.TraceField.DayOfYear: 100,
             segyio.TraceField.HourOfDay: 10,
             segyio.TraceField.MinuteOfHour: 30,
             segyio.TraceField.SecondOfMinute: 15}
        dt = C.datetime_from_header(h)
        # Independent re-derivation of "day-of-year 100" via ordinal
        # arithmetic (not the function's own timedelta path).
        expected_date = datetime.fromordinal(datetime(2024, 1, 1).toordinal() + 99)
        assert dt.date() == expected_date.date()
        assert (dt.hour, dt.minute, dt.second) == (10, 30, 15)

    def test_datetime_from_header_two_digit_year_offsets_to_2000s(self):
        h = {segyio.TraceField.YearDataRecorded: 24,
             segyio.TraceField.DayOfYear: 1,
             segyio.TraceField.HourOfDay: 0,
             segyio.TraceField.MinuteOfHour: 0,
             segyio.TraceField.SecondOfMinute: 0}
        dt = C.datetime_from_header(h)
        assert dt.year == 2024

    def test_datetime_from_header_zero_year_and_day_is_none(self):
        h = {segyio.TraceField.YearDataRecorded: 0,
             segyio.TraceField.DayOfYear: 0,
             segyio.TraceField.HourOfDay: 0,
             segyio.TraceField.MinuteOfHour: 0,
             segyio.TraceField.SecondOfMinute: 0}
        assert C.datetime_from_header(h) is None

    def test_coords_from_header_scalar_and_arcsec(self):
        h = {segyio.TraceField.SourceX: -28800,   # -8.0 deg * 3600
             segyio.TraceField.SourceY: 154800,    # 43.0 deg * 3600
             segyio.TraceField.SourceGroupScalar: 1,
             segyio.TraceField.CoordinateUnits: 2}
        x, y, cu = C.coords_from_header(h)
        assert x == pytest.approx(-8.0, abs=1e-9)
        assert y == pytest.approx(43.0, abs=1e-9)
        assert cu == 2

    def test_coords_from_header_zero_scalar_treated_as_one(self):
        h = {segyio.TraceField.SourceX: 500000,
             segyio.TraceField.SourceY: 4000000,
             segyio.TraceField.SourceGroupScalar: 0,
             segyio.TraceField.CoordinateUnits: 1}
        x, y, cu = C.coords_from_header(h)
        assert x == 500000.0 and y == 4000000.0 and cu == 1


# ════════════════════════════════════════════════════════════════════════════════
# Distance geodesy
# ════════════════════════════════════════════════════════════════════════════════

class TestDistance:
    def test_haversine_one_degree_latitude(self):
        # Independent re-derivation of the haversine formula (not calling
        # calc_distance's own code path) to catch a copy-paste regression.
        R = 6371000.0
        dphi = math.radians(1.0)
        a = math.sin(dphi / 2) ** 2
        c = 2 * math.atan2(math.sqrt(a), math.sqrt(1 - a))
        expected = R * c
        got = C.calc_distance(0.0, 0.0, 0.0, 1.0, cu=2)
        assert got == pytest.approx(expected, rel=1e-9)
        assert 111000 < got < 111400   # sanity band: ~111.2 km

    def test_haversine_same_point_is_zero(self):
        assert C.calc_distance(-8.0, 43.0, -8.0, 43.0, cu=2) == pytest.approx(0.0, abs=1e-6)

    def test_euclidean_3_4_5_triangle(self):
        assert C.calc_distance(0.0, 0.0, 3.0, 4.0, cu=1) == pytest.approx(5.0, abs=1e-9)

    def test_euclidean_zero_distance(self):
        assert C.calc_distance(100.0, 200.0, 100.0, 200.0, cu=1) == 0.0


# ════════════════════════════════════════════════════════════════════════════════
# UTM zone auto-detection
# ════════════════════════════════════════════════════════════════════════════════

class TestUtmZone:
    @pytest.mark.parametrize("lon,expected_zone", [
        (-180.0, 1), (-3.0, 30), (0.0, 31), (3.0, 31), (179.9, 60),
    ])
    def test_auto_zone(self, lon, expected_zone):
        assert C.auto_zone(lon) == expected_zone

    def test_to_utm_none_coords_pass_through(self):
        assert C.to_utm(None, 43.0) == (None, None, None)
        assert C.to_utm(-8.0, None) == (None, None, None)

    def test_to_utm_forced_zone_is_honoured(self):
        e, n, z = C.to_utm(-3.0, 43.0, force_zone=29)
        assert z == 29 and e is not None and n is not None

    def test_to_utm_auto_zone_matches_auto_zone_helper(self):
        e, n, z = C.to_utm(-3.0, 43.0)
        assert z == C.auto_zone(-3.0)

    def test_detect_zone_for_base_reads_first_sgy(self, tmp_path):
        from tests.make_synthetic_segy import make_synthetic_segy
        line_dir = tmp_path / "SGY" / "L1"
        line_dir.mkdir(parents=True)
        make_synthetic_segy(str(line_dir / "f.sgy"), n_traces=5, ns=16,
                            base_lon=-3.0, base_lat=43.0)
        zone, fname, lon = C.detect_zone_for_base(str(tmp_path))
        assert zone == 30
        assert fname == "f.sgy"
        assert lon == pytest.approx(-3.0, abs=1e-6)

    def test_detect_zone_for_base_no_sgy_files(self, tmp_path):
        (tmp_path / "SGY").mkdir()
        zone, fname, lon = C.detect_zone_for_base(str(tmp_path))
        assert zone is None and fname is None and lon is None

    def test_detect_zone_for_base_unreadable_coordinate(self, tmp_path):
        line_dir = tmp_path / "SGY" / "L1"
        line_dir.mkdir(parents=True)
        _write_metrics_sgy(str(line_dir / "z.sgy"), xs=[0, 0], ys=[0, 0], seconds=[0, 1])
        zone, fname, lon = C.detect_zone_for_base(str(tmp_path))
        assert zone is None and fname == "z.sgy" and lon is None


# ════════════════════════════════════════════════════════════════════════════════
# Acquisition statistics — ping rate / interval / vessel speed
# ════════════════════════════════════════════════════════════════════════════════

class TestAcquisitionMetrics:
    def test_single_file_exact_rate_interval_and_speed(self, tmp_path):
        # 11 traces spaced 2s apart (0..20s) -> delta_sec=20, rate=11/20=0.55 Hz,
        # interval=20/11 s. Straight line on the X axis, cu=1 (projected metres,
        # Euclidean distance) so speed is exactly deterministic: 300 m / 20 s.
        # xs starts at 100 (not 0) — nav_valida requires each endpoint to have
        # at least one nonzero coordinate; an all-zero start would be (wrongly
        # for this test) treated as "no navigation".
        line_dir = tmp_path / "SGY" / "L1"
        xs = [100.0 + i * 30.0 for i in range(11)]   # 100..400 in steps of 30
        ys = [0.0] * 11
        seconds = [i * 2 for i in range(11)]         # 0,2,...,20
        _write_metrics_sgy(str(line_dir / "f.sgy"), xs=xs, ys=ys, seconds=seconds)

        log, hz, s, kn = C.calculate_metrics_for_phase("FASE 1", str(tmp_path))
        assert len(hz) == 1 and len(s) == 1 and len(kn) == 1
        assert hz[0] == pytest.approx(11 / 20, rel=1e-9)
        assert s[0] == pytest.approx(20 / 11, rel=1e-9)
        expected_knots = (300.0 / 20.0) * 1.94384
        assert kn[0] == pytest.approx(expected_knots, rel=1e-9)
        assert "FASE 1" in log
        assert "MEDIA DE LA FASE" in log

    def test_averages_across_multiple_lines(self, tmp_path):
        # Two lines with different (but individually exact) rates; the
        # returned lists are per-FILE values, so the caller-side average is
        # just their mean — verify that arithmetic directly.
        base = tmp_path / "SGY"
        _write_metrics_sgy(str(base / "L1" / "a.sgy"),
                          xs=[0.0, 100.0], ys=[0.0, 0.0], seconds=[0, 10])
        _write_metrics_sgy(str(base / "L2" / "b.sgy"),
                          xs=[0.0, 50.0], ys=[0.0, 0.0], seconds=[0, 5])
        log, hz, s, kn = C.calculate_metrics_for_phase("FASE X", str(tmp_path))
        assert len(hz) == 2
        assert sorted(hz) == pytest.approx(sorted([2 / 10, 2 / 5]), rel=1e-9)
        avg_hz = sum(hz) / len(hz)
        assert avg_hz == pytest.approx((2 / 10 + 2 / 5) / 2, rel=1e-9)

    def test_insufficient_traces_is_skipped(self, tmp_path):
        line_dir = tmp_path / "SGY" / "L1"
        _write_metrics_sgy(str(line_dir / "one.sgy"), xs=[0.0], ys=[0.0], seconds=[0])
        log, hz, s, kn = C.calculate_metrics_for_phase("FASE 1", str(tmp_path))
        assert hz == [] and s == [] and kn == []
        assert "Insuficientes trazas" in log
        assert "No se pudo calcular ninguna métrica" in log

    def test_invalid_dates_are_skipped(self, tmp_path):
        line_dir = tmp_path / "SGY" / "L1"
        _write_metrics_sgy(str(line_dir / "bad.sgy"), xs=[0.0, 10.0], ys=[0.0, 0.0],
                          seconds=[0, 1], year=0, doy=0)
        log, hz, s, kn = C.calculate_metrics_for_phase("FASE 1", str(tmp_path))
        assert hz == [] and s == []
        assert "Fechas a cero o inválidas" in log

    def test_zero_delta_time_is_skipped(self, tmp_path):
        line_dir = tmp_path / "SGY" / "L1"
        _write_metrics_sgy(str(line_dir / "same.sgy"), xs=[0.0, 10.0], ys=[0.0, 0.0],
                          seconds=[5, 5])   # identical timestamps
        log, hz, s, kn = C.calculate_metrics_for_phase("FASE 1", str(tmp_path))
        assert hz == [] and s == []
        assert "Delta de tiempo es 0s" in log

    def test_missing_navigation_still_yields_rate_but_no_speed(self, tmp_path):
        line_dir = tmp_path / "SGY" / "L1"
        # Valid time delta, but both endpoints at (0,0) -> nav_valida is False.
        _write_metrics_sgy(str(line_dir / "nonav.sgy"), xs=[0.0, 0.0], ys=[0.0, 0.0],
                          seconds=[0, 10])
        log, hz, s, kn = C.calculate_metrics_for_phase("FASE 1", str(tmp_path))
        assert len(hz) == 1 and len(s) == 1
        assert kn == []
        assert "V: N/A" in log

    def test_empty_directory_reports_no_metrics(self, tmp_path):
        (tmp_path / "SGY").mkdir()
        log, hz, s, kn = C.calculate_metrics_for_phase("FASE VACIA", str(tmp_path))
        assert hz == [] and s == [] and kn == []
        assert "No se pudo calcular ninguna métrica" in log

    def test_flat_directory_without_sgy_subfolder(self, tmp_path):
        # _resolve_sgy_base falls back to base_dir itself when SGY/ is absent.
        _write_metrics_sgy(str(tmp_path / "L1" / "f.sgy"),
                          xs=[0.0, 20.0], ys=[0.0, 0.0], seconds=[0, 4])
        log, hz, s, kn = C.calculate_metrics_for_phase("FASE 1", str(tmp_path))
        assert len(hz) == 1
        assert hz[0] == pytest.approx(2 / 4, rel=1e-9)


# ════════════════════════════════════════════════════════════════════════════════
# Excel builders
# ════════════════════════════════════════════════════════════════════════════════

class TestExcelBuilders:
    def _make_project(self, tmp_path, n_lines=2):
        from tests.make_synthetic_segy import make_synthetic_segy
        base = tmp_path / "proj"
        for i in range(1, n_lines + 1):
            d = base / "SGY" / f"L{i}"
            d.mkdir(parents=True)
            make_synthetic_segy(str(d / f"2024010112000{i}_a.sgy"),
                                n_traces=20, ns=32, base_lon=-8.0 + i * 0.01,
                                base_lat=43.0, lon_step=0.001)
        return base

    def test_build_registro_excel_creates_expected_sheet_and_rows(self, tmp_path):
        base = self._make_project(tmp_path, n_lines=2)
        out = str(tmp_path / "reg.xlsx")
        summary = C.build_registro_excel(
            [{"sheet_name": "1ª FASE", "label": "FASE UNO", "base_dir": str(base)}],
            "PROYECTO_TEST", out)
        assert Path(out).exists()
        assert len(summary) == 1 and "2 l" in summary[0]

        wb = openpyxl.load_workbook(out)
        assert "1ª FASE" in wb.sheetnames
        ws = wb["1ª FASE"]
        assert ws["A1"].value == "PROYECTO_TEST"
        assert ws["A4"].value == "FASE UNO"
        # Two line rows starting at row 5 (L1, L2), each with a filename in
        # the "FICHERO INICIO (*.sgy)" column (col 8).
        assert ws.cell(row=5, column=1).value in ("L1", "L2")
        assert ws.cell(row=6, column=1).value in ("L1", "L2")
        assert str(ws.cell(row=5, column=8).value).endswith(".sgy")

    def test_build_registro_excel_multi_phase(self, tmp_path):
        base1 = self._make_project(tmp_path / "p1", n_lines=1)
        base2 = self._make_project(tmp_path / "p2", n_lines=3)
        out = str(tmp_path / "reg2.xlsx")
        summary = C.build_registro_excel(
            [{"sheet_name": "FASE A", "label": "A", "base_dir": str(base1)},
             {"sheet_name": "FASE B", "label": "B", "base_dir": str(base2)}],
            "PROY", out)
        assert len(summary) == 2
        wb = openpyxl.load_workbook(out)
        assert set(wb.sheetnames) == {"FASE A", "FASE B"}

    def test_build_coordenadas_excel_has_both_sheets_and_formulas(self, tmp_path):
        base = self._make_project(tmp_path, n_lines=2)
        out = str(tmp_path / "coord.xlsx")
        summary = C.build_coordenadas_excel(
            [{"label": "FASE UNO", "base_dir": str(base), "forced_zone": None}],
            "PROY", out)
        assert Path(out).exists()
        assert len(summary) == 1
        assert "con coordenadas" in summary[0]

        wb = openpyxl.load_workbook(out)
        assert "COORDENADAS LINEAS" in wb.sheetnames
        assert "LONGITUD LINEAS" in wb.sheetnames
        ws_long = wb["LONGITUD LINEAS"]
        # Column headers occupy row 1, so the phase section starts at row 2.
        # The section label gets a "(HUSO n)" suffix appended once the zone(s)
        # actually used are known — see the exact-format test below.
        assert ws_long.cell(row=2, column=1).value.startswith("FASE UNO (HUSO")
        # First data row (row 3) should carry the length formulas.
        formula_m = ws_long.cell(row=3, column=10).value
        assert isinstance(formula_m, str) and formula_m.startswith("=SQRT(")
        assert ws_long.cell(row=3, column=11).value == "=J3/1000"
        assert ws_long.cell(row=3, column=12).value == "=J3/1852"

    def test_build_coordenadas_excel_forced_zone_marker(self, tmp_path):
        base = self._make_project(tmp_path, n_lines=1)
        out = str(tmp_path / "coord_forced.xlsx")
        C.build_coordenadas_excel(
            [{"label": "FASE UNO", "base_dir": str(base), "forced_zone": 30}],
            "PROY", out)
        wb = openpyxl.load_workbook(out)
        ws = wb["COORDENADAS LINEAS"]
        label = ws.cell(row=2, column=1).value
        assert "HUSO 30" in label
        assert label == "FASE UNO (HUSO 30 *)"   # exact forced-zone format

    def test_build_coordenadas_excel_no_sgy_data(self, tmp_path):
        base = tmp_path / "empty"
        (base / "SGY" / "L1").mkdir(parents=True)
        out = str(tmp_path / "coord_empty.xlsx")
        summary = C.build_coordenadas_excel(
            [{"label": "SIN DATOS", "base_dir": str(base), "forced_zone": None}],
            "PROY", out)
        assert "sin coordenadas SGY" in summary[0]
